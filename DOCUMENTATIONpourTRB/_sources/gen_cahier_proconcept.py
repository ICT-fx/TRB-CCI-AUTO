# -*- coding: utf-8 -*-
"""Cahier des charges ProConcept — court, adressé à la personne ProConcept."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docgen_common import *
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "Cahier_des_charges_ProConcept.docx")

doc = new_doc()

# ============ PAGE DE TITRE ============
title_block(doc, [
    ("CAHIER DES CHARGES PROCONCEPT", {"bold": True, "color": TRB_BLUE, "size": 23, "space_after": 2, "space_before": 34}),
    ("Import des commandes clients (CCI) dans ProConcept", {"bold": True, "color": DARK, "size": 15, "space_after": 10}),
    ("Ce que nous fournissons · ce que ProConcept doit compléter · le livrable attendu de votre part",
     {"italic": True, "color": GREY, "size": 11.5, "space_after": 16}),
])
kv_block(doc, [
    ("Destinataire", "Responsable / intégrateur ProConcept"),
    ("Émetteur", "TRB Chemedica International SA — Supply Chain"),
    ("Objet", "Définir le format d'import permettant de créer automatiquement les CCI"),
    ("Version", "1.0  —  3 juillet 2026"),
])
para(doc, "")
callout(doc,
        "À partir des quelques champs listés ci-dessous, nous voulons créer une commande client "
        "interne (CCI) dans ProConcept. Toutes les autres informations doivent être déduites "
        "automatiquement du code client (Clé 1) et du code article (SKU). Nous attendons de vous "
        "un gabarit d'import (Excel, CSV ou le format de votre choix) qui charge ces données et "
        "crée les CCI — testé et fonctionnel.",
        label="En une phrase :", fill=INFO_FILL, edge="548235")

# ============ 1. CONTEXTE ============
doc.add_heading("1. Contexte et objectif", level=1)
para(doc, "TRB Chemedica reçoit ses commandes clients sous des formes hétérogènes (PDF, scans, "
     "images). Un outil interne les lit automatiquement et en extrait des données propres et "
     "normalisées, déjà rapprochées de nos référentiels : le code client (« Clé 1 ») et le code "
     "article (SKU) sont résolus et validés en amont.")
para(doc, "Nous voulons charger ces données dans ProConcept pour créer les CCI sans ressaisie "
     "manuelle. L'objet de ce document est de définir clairement :")
bullet(doc, [("les données que nous vous transmettons", True), (" (section 2) ;", False)])
bullet(doc, [("les données que ProConcept doit compléter automatiquement", True), (" à partir de la Clé 1 et du SKU (section 3) ;", False)])
bullet(doc, [("le livrable attendu de votre part", True), (" : le gabarit d'import (section 4).", False)])

# ============ 2. DONNEES FOURNIES ============
doc.add_heading("2. Les données que nous fournissons", level=1)
para(doc, "Un fichier d'import correspond à un lot : il regroupe plusieurs commandes. Chaque "
     "commande (CCI) a un en-tête (le client et ses références) et une ou plusieurs lignes "
     "(article + quantité). Nous pouvons produire ces champs dans l'ordre, avec les intitulés et "
     "le format que vous préciserez dans votre gabarit.")

para(doc, "En-tête de commande (une fois par CCI)", bold=True, color=TRB_BLUE, space_before=4)
table(doc,
      ["Champ", "Description", "Rôle"],
      [
        ["Clé 1", "Code client dans ProConcept", "Obligatoire — identifie le client"],
        ["Nom du client", "Nom du client (lisible)", "Informatif — peut être ignoré par le système"],
        ["Référence partenaire", "Référence de la commande côté client (n° du bon de commande)", "Obligatoire — référence externe de la commande"],
        ["Date de livraison souhaitée", "Date de livraison demandée", "Obligatoire"],
      ],
      widths=[1.9, 3.0, 1.8])

para(doc, "Lignes de commande (une ou plusieurs par CCI)", bold=True, color=TRB_BLUE, space_before=4)
table(doc,
      ["Champ", "Description", "Rôle"],
      [
        ["SKU", "Code article TRB (4 chiffres)", "Obligatoire — identifie l'article"],
        ["Quantité", "Quantité commandée", "Obligatoire"],
      ],
      widths=[1.9, 3.0, 1.8])

callout(doc,
        "Un fichier d'import = un LOT de plusieurs commandes (pas un fichier par commande). "
        "Au sein du lot, les lignes d'une même commande sont regroupées par leur en-tête — "
        "nous prévoyons la clé « Clé 1 + référence partenaire ».",
        label="Important :")

para(doc, "Le nom du client n'est là que pour faciliter le contrôle visuel : il peut être ignoré "
     "à l'import. Le format exact (colonnes, ordre, format de date, séparateur décimal, encodage) "
     "suivra votre gabarit — voir sections 4 et 5.", italic=True, color=GREY, size=10)

# ============ 3. A COMPLETER PAR PROCONCEPT ============
doc.add_heading("3. Ce que ProConcept doit compléter automatiquement", level=1)
para(doc, "Nous ne transmettons volontairement que les champs de la section 2. Toutes les autres "
     "informations nécessaires à la commande doivent être déduites automatiquement des données "
     "déjà présentes dans ProConcept :")
bullet(doc, [("À partir du client (Clé 1) : ", True),
             ("monnaie, conditions de paiement, incoterm, adresse de livraison par défaut, mode "
              "d'expédition, régime de TVA, tarif / liste de prix applicable, langue, etc.", False)])
bullet(doc, [("À partir de l'article (SKU) : ", True),
             ("désignation, prix unitaire (selon le tarif du client), unité de vente / "
              "conditionnement, etc.", False)])
callout(doc,
        "Principe : nous fournissons l'identité de la commande (qui, quoi, combien, quand, quelle "
        "référence) ; ProConcept fournit tout le reste à partir des fiches client et article.",
        label="Répartition :")

# ============ 4. LIVRABLE ATTENDU ============
doc.add_heading("4. Ce que nous attendons de vous (le livrable)", level=1)
para(doc, "Un gabarit / masque d'import — au format de votre choix (Excel, CSV, XML… tout format "
     "supporté par ProConcept) — qui nous permette de charger nos commandes et de créer les CCI. "
     "Concrètement, nous attendons :")
bullet(doc, [("Le gabarit lui-même", True), (" : colonnes attendues (intitulés exacts, ordre, champs obligatoires / optionnels).", False)])
bullet(doc, [("Les règles de format", True), (" : format de date, séparateur décimal, encodage, séparateur CSV, etc.", False)])
bullet(doc, [("La procédure de chargement", True), (" : comment importer le fichier dans ProConcept pour créer les CCI (fonction d'import ou étapes à suivre).", False)])
bullet(doc, [("La déduction automatique", True), (" : confirmation que les champs de la section 3 sont bien complétés à partir de la Clé 1 et du SKU.", False)])
bullet(doc, [("Un test concluant", True), (" : la démonstration que l'import fonctionne sur un jeu d'exemple que nous fournirons.", False)])
para(doc, "Idéalement : un fichier gabarit vide, un exemple rempli, et une courte note de procédure.",
     italic=True, color=GREY, size=10)

# ============ 5. POINTS A PRECISER ============
doc.add_heading("5. Points à préciser ensemble", level=1)
para(doc, "Quelques questions pour caler le gabarit :")
bullet(doc, "Quels champs sont réellement obligatoires côté ProConcept pour créer une CCI ? Notre liste (section 2) est-elle suffisante, ou manque-t-il quelque chose ?")
bullet(doc, "Un fichier contient un lot de plusieurs commandes : confirmez-vous que ProConcept sait créer plusieurs CCI depuis un seul fichier, et que le regroupement des lignes par « Clé 1 + référence partenaire » vous convient ?")
bullet(doc, "Existe-t-il une fonction d'import standard dans ProConcept (Excel / CSV / XML), ou faut-il un développement spécifique ?")
bullet(doc, "Quels format de date, séparateur décimal et encodage attendez-vous ?")
bullet(doc, "Que se passe-t-il si un client (Clé 1) ou un article (SKU) n'existe pas encore dans ProConcept — rejet, ou création ?")
bullet(doc, "Le prix unitaire doit-il venir du tarif ProConcept, ou pouvons-nous le laisser vide ?")

# ============ 6. EXEMPLE ============
doc.add_heading("6. Exemple de jeu de données (illustratif)", level=1)
para(doc, "Un seul fichier contenant un lot de deux commandes. La mise en forme exacte (colonnes, "
     "ordre) suivra votre gabarit — cet exemple sert uniquement à illustrer le contenu.")
table(doc,
      ["Clé 1", "Nom du client (ignoré)", "Référence partenaire", "Date de livraison", "SKU", "Quantité"],
      [
        ["2780008", "Pharmacie Dupont", "CMD-4471", "15/07/2026", "1311", "100"],
        ["2780008", "Pharmacie Dupont", "CMD-4471", "15/07/2026", "1274", "50"],
        ["2455301", "Clinique du Lac", "BC-2098", "20/07/2026", "1311", "30"],
      ],
      widths=[0.9, 1.7, 1.5, 1.3, 0.7, 0.8])
para(doc, "Ce fichier est un lot de deux commandes : la première (Clé 1 2780008 / réf. CMD-4471) a "
     "deux lignes, la seconde (Clé 1 2455301 / réf. BC-2098) en a une. Les lignes sont regroupées "
     "en commandes par « Clé 1 + référence partenaire » → ProConcept crée ici deux CCI.",
     italic=True, color=GREY, size=10)

doc.save(OUT)
print("OK:", OUT)
