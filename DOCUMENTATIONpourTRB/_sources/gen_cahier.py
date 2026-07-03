# -*- coding: utf-8 -*-
"""Cahier des charges — Outil CCI (TRB Chemedica)."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docgen_common import *
from docx.enum.text import WD_ALIGN_PARAGRAPH

FLOW = os.path.join(os.path.dirname(os.path.abspath(__file__)), "flowchart_cci.png")
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), r"Cahier_des_charges_Outil_CCI_v1.docx")

doc = new_doc()

# ============ PAGE DE TITRE ============
title_block(doc, [
    ("CAHIER DES CHARGES", {"bold": True, "color": TRB_BLUE, "size": 24, "space_after": 2, "space_before": 40}),
    ("Outil CCI — Extraction automatique des commandes clients", {"bold": True, "color": DARK, "size": 15, "space_after": 2}),
    ("et pré-saisie ProConcept", {"bold": True, "color": DARK, "size": 15, "space_after": 12}),
    ("Spécification fonctionnelle de l'outil en production (Azure Function + orchestration SharePoint)",
     {"italic": True, "color": GREY, "size": 11.5, "space_after": 18}),
])
para(doc, "", space_after=2)
info = [
    ("Organisation", "TRB Chemedica International SA — Supply Chain"),
    ("Destinataires", "Équipe projet / Supply Chain / IT"),
    ("Objet", "Décrire ce que fait l'outil, ses règles, et le rôle de l'humain dans le processus"),
    ("Version", "1.0"),
    ("Date", "3 juillet 2026"),
    ("Statut", "Documentation de l'outil déployé — base de cadrage et d'évolution"),
]
kv_block(doc, info)
para(doc, "")
para(doc, "Ce document décrit l'outil « CCI » réalisé et déployé par l'équipe (Azure Function "
     "+ IA Claude + orchestration SharePoint). Il partage l'objectif de la démarche menée "
     "en parallèle sous Microsoft Copilot / Cowork, mais en couvre un périmètre plus ciblé "
     "(voir § 1.3). Un guide utilisateur distinct accompagne ce cahier des charges.",
     italic=True, color=GREY, size=10)

doc.add_page_break()

# ============ SOMMAIRE ============
doc.add_heading("Sommaire", level=1)
sommaire(doc, [
    ("1. Contexte et objectif", False),
    ("1.1 Contexte  ·  1.2 Objectif  ·  1.3 Positionnement", True),
    ("2. Périmètre", False),
    ("2.1 Inclus  ·  2.2 Exclus / non couvert à ce jour", True),
    ("3. Acteurs et déclenchement", False),
    ("4. Architecture de l'outil", False),
    ("5. Exigences fonctionnelles", False),
    ("6. Données et référentiel (master data)", False),
    ("6.1 Master data  ·  6.2 Provenance & rafraîchissement  ·  6.3 Gouvernance", True),
    ("7. Règles de gestion", False),
    ("8. Exigences non fonctionnelles", False),
    ("9. Scénario nominal et cas limites", False),
    ("10. Critères d'acceptation", False),
    ("11. Hypothèses et dépendances", False),
    ("12. Hors périmètre et évolutions prévues", False),
    ("13. Glossaire", False),
])

# ============ 1. CONTEXTE ============
doc.add_heading("1. Contexte et objectif", level=1)
doc.add_heading("1.1 Contexte", level=2)
para(doc, "TRB Chemedica reçoit ses commandes clients (distributeurs et filiales) par e-mail, "
     "sous forme de bons de commande hétérogènes (PDF natifs ou scannés, images, y compris "
     "documents manuscrits). Leur saisie dans l'ERP ProConcept est aujourd'hui manuelle, "
     "chronophage et exposée aux erreurs — notamment sur le code client et les codes article (SKU).")
para(doc, "Pour fiabiliser et accélérer cette saisie, un outil dédié a été développé et déployé : "
     "il lit chaque bon de commande avec l'IA (Claude), retrouve le code client (« Clé 1 ») et "
     "confirme / corrige les SKU à partir de l'historique des ventes, puis produit un fichier "
     "Excel de pré-saisie prêt à importer dans ProConcept. Les commandes qui ne peuvent pas "
     "être traitées automatiquement sont routées vers une revue humaine.")

doc.add_heading("1.2 Objectif", level=2)
bullet(doc, "Réduire le temps de saisie des commandes et fiabiliser les données avant import dans ProConcept.")
bullet(doc, [("Retrouver automatiquement le ", False), ("code client (Clé 1)", True),
             (" à partir du nom, et ", False), ("confirmer ou corriger le SKU", True),
             (" de chaque ligne via le catalogue réellement acheté par ce client.", False)])
bullet(doc, [("Garder l'humain dans la boucle", True),
             (" : l'outil prépare et signale ses doutes ; le gestionnaire vérifie et valide. "
              "Rien n'est écrit dans ProConcept sans contrôle humain.", False)])

doc.add_heading("1.3 Positionnement (par rapport à la démarche Copilot / Cowork)", level=2)
para(doc, "Un agent plus large a été prototypé en parallèle sous Microsoft Copilot / Cowork "
     "(cycle e-mail complet, brouillon de confirmation, génération XML ProConcept, etc.). "
     "L'outil décrit ici poursuit le même but mais avec un périmètre volontairement plus "
     "ciblé et déjà en production. Les principales différences :")
table(doc,
      ["Aspect", "Outil CCI (ce document)", "Agent Copilot / Cowork"],
      [
        ["Réception des e-mails",
         "Règle de messagerie simple (certains expéditeurs + pièce jointe) puis tri humain",
         "Détection sémantique de chaque e-mail (commande / doute / non-commande)"],
        ["Lecture du document", "IA (Claude) — PDF natif/scanné, images, manuscrit", "IA + OCR multi-format (Excel, Word…)"],
        ["Résolution", "Clé 1 + SKU via catalogue client", "Clé 1 + SKU + détection anciens codes (Tyvek)"],
        ["Sortie", "1 Excel CCI consolidé (1 feuille, 1 ligne/commande)", "1 Excel à 2 onglets (Vérification + Load) + XML"],
        ["Confirmation client", "Hors périmètre (aucun e-mail)", "Brouillon de confirmation préparé (non envoyé)"],
        ["Import ProConcept", "Manuel par le gestionnaire", "Manuel (fichiers de chargement générés)"],
        ["Socle technique", "Azure Function + orchestration SharePoint", "Microsoft 365 / Power Platform / Copilot"],
      ],
      widths=[1.5, 2.7, 2.5])
para(doc, "Les deux approches convergent sur le principe fondamental : l'outil prépare, "
     "l'humain valide, l'ERP n'est jamais alimenté automatiquement.", italic=True, color=GREY, size=10)

# ============ 2. PERIMETRE ============
doc.add_heading("2. Périmètre", level=1)
doc.add_heading("2.1 Inclus", level=2)
bullet(doc, "Dépôt automatique des pièces jointes de certains expéditeurs connus dans le dossier d'entrée (règle de messagerie), complété par un tri humain.")
bullet(doc, "Traitement par lot des commandes déposées dans le dossier d'entrée « Commandes-PDF » (quotidien à heure fixe + lancement manuel).")
bullet(doc, "Lecture par l'IA de chaque bon de commande (PDF natif, PDF scanné, image PNG/JPEG/TIFF, y compris manuscrit).")
bullet(doc, "Extraction structurée : nom du client, référence partenaire (n° de commande client), date de livraison souhaitée ; par ligne : désignation, SKU, quantité.")
bullet(doc, "Résolution du numéro client (« Clé 1 ») par correspondance du nom avec le master data.")
bullet(doc, "Confirmation / correction du SKU de chaque ligne via le catalogue réellement acheté par ce client.")
bullet(doc, "Routage automatique : commande traitée → « Commandes-Done » ; commande à problème → « A-revoir ».")
bullet(doc, "Renommage des documents traités sous un nom lisible « Client - JJ-MM-AAAA ».")
bullet(doc, "Génération d'UN SEUL fichier Excel CCI consolidé par lot (1 ligne par commande), horodaté, déposé dans « Uppload-CCI ».")
bullet(doc, "Génération d'un Excel de reporting des commandes rejetées (dossier « A-revoir »).")
bullet(doc, "Surlignage jaune des points incertains (SKU ambigu, Clé 1 non résolue) pour orienter la vérification humaine.")

doc.add_heading("2.2 Exclus / non couvert à ce jour", level=2)
bullet(doc, "Détection sémantique des e-mails de commande (l'entrée repose sur une règle d'expéditeur + un tri humain).")
bullet(doc, "Envoi ou brouillon de confirmation au client (aucun e-mail n'est produit ni envoyé).")
bullet(doc, "Écriture / intégration automatique directe dans ProConcept (l'outil produit un fichier de pré-saisie ; l'import reste manuel).")
bullet(doc, "Génération d'un fichier XML de chargement ProConcept et détection spécifique des « anciens codes » (Tyvek).")
bullet(doc, "Extraction du prix, de la monnaie, de l'incoterm, des conditions de paiement et de l'adresse de facturation (non repris dans l'Excel CCI).")
bullet(doc, "Reconstruction ou maintenance de la base client de référence (fournie et rafraîchie hors outil).")

# ============ 3. ACTEURS ============
doc.add_heading("3. Acteurs et déclenchement", level=1)
bullet(doc, [("Acteur principal : ", True), ("gestionnaire / spécialiste Supply Chain", False),
             (" — dépose/trie les commandes, vérifie l'Excel de pré-saisie, saisit dans ProConcept.", False)])
bullet(doc, [("Déclenchement planifié : ", True), ("traitement par lot chaque jour à 15 h", False),
             (" sur le contenu du dossier « Commandes-PDF ».", False)])
bullet(doc, [("Déclenchement à la demande : ", True), ("lancement manuel du lot", False),
             (" à tout moment (ex. commande urgente).", False)])
bullet(doc, [("Systèmes connectés : ", True),
             ("SharePoint (dossiers d'entrée/sortie/archive), Azure Function d'extraction, "
              "IA Claude, master data (historique des ventes), et — en sortie — l'Excel CCI importé dans ProConcept.", False)])

# ============ 4. ARCHITECTURE ============
doc.add_heading("4. Architecture de l'outil", level=1)
para(doc, "L'outil repose sur une fonction serverless (Azure Function, Python) qui porte toute la "
     "logique métier, orchestrée par un flux qui parcourt le dossier d'entrée et agrège les "
     "résultats. Le traitement est unitaire (un document par appel), ce qui isole les échecs : "
     "un mauvais scan ne bloque pas le reste du lot.")
table(doc,
      ["Brique", "Rôle", "Notes"],
      [
        ["Dossiers SharePoint", "Entrée / sortie / archives : Commandes-PDF, Uppload-CCI, Commandes-Done, A-revoir",
         "Bibliothèque documentaire dédiée ; alimentée par la règle mail + le tri humain."],
        ["Orchestration (flux)", "Parcourt Commandes-PDF, appelle l'extraction pour chaque fichier, agrège les lignes, range les fichiers",
         "Lot séquentiel, quotidien + manuel ; déplace le PDF selon le résultat (Done / A-revoir)."],
        ["Azure Function — /extract", "Lit 1 document, extrait les données, résout le master data, renvoie une ligne (ou un rejet)",
         "Cœur du traitement ; 2 appels IA (extraction + résolution)."],
        ["Azure Function — /build", "Agrège toutes les lignes du lot en UN Excel CCI consolidé",
         "1 feuille, 1 ligne par commande ; horodaté, jamais d'écrasement."],
        ["Azure Function — /build_errors", "Produit l'Excel de reporting des commandes rejetées",
         "Colonnes : fichier · date · observation."],
        ["IA (Claude)", "Lecture du document (vision) + résolution master data",
         "Appels par outil forcé (sortie typée) ; garde-fou anti-invention (§ 7)."],
        ["Master data", "Référentiel clients (Clé 1) + catalogue par client",
         "Dérivé de l'historique des ventes ; embarqué avec la fonction, rafraîchi périodiquement."],
      ],
      widths=[1.7, 2.7, 2.3])

# ============ 5. EXIGENCES FONCTIONNELLES ============
doc.add_heading("5. Exigences fonctionnelles", level=1)
req(doc, "EF-1", "Dépôt automatique (règle de messagerie)",
    "les pièces jointes des e-mails provenant de certains expéditeurs connus sont déposées "
    "automatiquement dans l'espace de réception. Cette règle est un filet de départ, pas une "
    "détection sémantique.")
req(doc, "EF-2", "Tri humain de la réception",
    "le gestionnaire vérifie les documents reçus et dépose chaque commande à traiter dans le "
    "dossier « Commandes-PDF ». Les cas non captés par la règle (commande dans le corps du mail, "
    "expéditeur inconnu, pièce jointe non pertinente) sont traités à la main.")
req(doc, "EF-3", "Traitement par lot",
    "à heure fixe (15 h) et sur lancement manuel, l'outil traite tous les documents présents "
    "dans « Commandes-PDF », un par un.")
req(doc, "EF-4", "Lecture multi-format",
    "l'outil lit les PDF natifs, les PDF scannés et les images (PNG, JPEG, TIFF), y compris les "
    "documents manuscrits ou avec valeurs corrigées à la main. Les TIFF sont convertis avant lecture.")
req(doc, "EF-5", "Extraction structurée",
    "l'outil extrait le nom du client, la référence partenaire (n° de commande client), la date de "
    "livraison souhaitée ; et, par ligne : la désignation, le SKU tel qu'écrit, la quantité. "
    "L'adresse de livraison est lue uniquement pour distinguer des clients homonymes (ex. TRB "
    "Thailand / Vietnam / Myanmar) et n'apparaît pas dans l'Excel.")
req(doc, "EF-6", "Validation de format",
    "avant résolution, une commande illisible ou incomplète (nom de client, référence ou date "
    "manquants, aucune ligne, quantité ≤ 0) est rejetée vers « A-revoir ». Le SKU n'est pas "
    "contrôlé à ce stade : il sera corrigé via le catalogue.")
req(doc, "EF-7", "Résolution du numéro client (Clé 1)",
    "l'outil rapproche le nom lu de la table clients du master data et renvoie la « Clé 1 », en "
    "gérant les variantes de dénomination (casse, forme juridique, abréviations).")
req(doc, "EF-8", "Confirmation / correction du SKU",
    "pour chaque ligne, l'outil compare le SKU et la désignation au catalogue de CE client : le "
    "SKU est confirmé s'il est correct, ou remplacé par le bon SKU retrouvé via la désignation.")
req(doc, "EF-9", "Statut de résolution par ligne",
    "chaque ligne reçoit un statut interne : correct, corrigé, ambigu (l'IA hésite entre plusieurs "
    "produits proches), ou inconnu (ni le SKU ni la désignation ne matchent le catalogue du client).")
req(doc, "EF-10", "Rejet strict après résolution",
    "si le client est introuvable dans le master data, ou si au moins un produit est hors du "
    "catalogue de ce client, la commande est routée vers « A-revoir » plutôt que de produire une "
    "ligne fausse.")
req(doc, "EF-11", "Routage et renommage",
    "une commande traitée voit son PDF renommé « Client - JJ-MM-AAAA » et déplacé dans "
    "« Commandes-Done » ; une commande rejetée est renommée de la même façon et déplacée dans "
    "« A-revoir ».")
req(doc, "EF-12", "Excel CCI consolidé unique",
    "toutes les commandes valides d'un lot sont regroupées dans UN SEUL fichier Excel horodaté "
    "(ex. CCI-Lot-AAAAMMJJ-HHMMSS.xlsx), 1 ligne par commande, déposé dans « Uppload-CCI » ; "
    "jamais un fichier par commande, jamais d'écrasement.")
req(doc, "EF-13", "Colonnes de l'Excel CCI",
    "Nom du client · Clé 1 · Référence partenaire · Date de livraison souhaitée · puis, par produit, "
    "SKU i / Quantité i (jusqu'au maximum du lot) · et Nom du fichier (nom composé, pour retrouver "
    "la commande d'origine).")
req(doc, "EF-14", "Surlignage des incertitudes",
    "les cellules à contrôler en priorité sont surlignées en jaune : un SKU ambigu (l'IA n'est pas "
    "sûre) et une Clé 1 vide (anormale). Les corrections jugées évidentes ne sont pas surlignées.")
req(doc, "EF-15", "Rapport des commandes rejetées",
    "l'outil produit un Excel de reporting listant chaque commande envoyée en « A-revoir », avec "
    "trois colonnes : Nom du fichier · Date de l'erreur · Observation (motif du rejet).")
req(doc, "EF-16", "Neutralisation des formules",
    "toute valeur issue d'un document client susceptible d'être interprétée comme une formule "
    "Excel est neutralisée à l'écriture (protection contre l'injection de formules).")

# ============ 6. DONNEES ============
doc.add_heading("6. Données et référentiel (master data)", level=1)
doc.add_heading("6.1 Master data (clients + catalogue)", level=2)
para(doc, "La résolution s'appuie sur un référentiel dérivé de l'historique des ventes, rangé en deux tables :")
bullet(doc, [("Table clients", True), (" : à chaque nom de client correspond sa « Clé 1 » (numéro client, 7 chiffres).", False)])
bullet(doc, [("Catalogue par client", True), (" : pour chaque client, la liste des produits qu'il achète "
             "réellement (SKU à 4 chiffres + désignation), soit une ligne par couple client-SKU.", False)])
para(doc, "Ce référentiel suffit : il fournit le numéro client, l'historique des SKU achetés par chaque "
     "client (base de la confirmation/correction) et les désignations permettant de retrouver le bon SKU.")

doc.add_heading("6.2 Provenance et rafraîchissement", level=2)
bullet(doc, "Le master data est généré depuis un export de l'historique des ventes (déduplication automatique) et embarqué avec l'outil.")
bullet(doc, "Il doit être rafraîchi périodiquement : plus il est à jour, plus l'outil résout de clients et de SKU automatiquement (moins de « A-revoir »).")
bullet(doc, "Un nouvel export met à jour aussi bien la liste des clients que les catalogues (nouveaux produits achetés).")

doc.add_heading("6.3 Gouvernance de la donnée", level=2)
bullet(doc, "Désigner un propriétaire du référentiel et une fréquence de rafraîchissement (ex. mensuelle).")
bullet(doc, "Un client absent du master data ou un produit jamais acheté par ce client provoque un « A-revoir » légitime : c'est le signal d'un référentiel à compléter (nouveau client / nouveau produit).")

# ============ 7. REGLES DE GESTION ============
doc.add_heading("7. Règles de gestion", level=1)
req(doc, "RG-1", "Source unique", "la résolution s'appuie exclusivement sur le master data (clients + catalogue) dérivé de l'historique des ventes.")
req(doc, "RG-2", "Document client uniquement", "la commande est lue depuis le document du client, jamais reconstituée depuis une facture ou une confirmation TRB.")
req(doc, "RG-3", "Ne rien inventer", "un code client ou un SKU n'est retenu que s'il existe réellement dans le master data ; aucune valeur n'est fabriquée (garde-fou anti-hallucination).")
req(doc, "RG-4", "Doute = signalé", "un SKU pour lequel l'IA hésite est marqué « ambigu » et surligné jaune ; il n'est jamais imposé silencieusement.")
req(doc, "RG-5", "Rejet plutôt que faux", "client introuvable ou produit hors catalogue → « A-revoir » ; l'outil ne produit pas de ligne incertaine.")
req(doc, "RG-6", "Fichier neuf à chaque lot", "un Excel CCI horodaté est créé à chaque exécution ; jamais d'écrasement.")
req(doc, "RG-7", "Fichier unique par lot", "toutes les commandes valides d'un même lot sont regroupées dans un seul Excel de pré-saisie.")
req(doc, "RG-8", "Humain dans la boucle", "la pré-saisie est vérifiée par le gestionnaire avant tout import dans ProConcept ; l'outil n'écrit jamais dans l'ERP.")
req(doc, "RG-9", "Traçabilité du nom", "chaque commande traitée ou rejetée est renommée « Client - JJ-MM-AAAA », nom repris dans l'Excel pour faire le lien avec le document d'origine.")
req(doc, "RG-10", "Isolation des échecs", "un document illisible ou en erreur ne bloque pas le lot : il part en « A-revoir » et le traitement continue.")

# ============ 8. EXIGENCES NON FONCTIONNELLES ============
doc.add_heading("8. Exigences non fonctionnelles", level=1)
req(doc, "ENF-1", "Sécurité et permissions", "principe du moindre privilège : lecture/écriture limitée aux dossiers de l'outil ; clé d'API stockée côté serveur (jamais dans le code) ; aucune écriture dans ProConcept.")
req(doc, "ENF-2", "Jamais écrire dans ProConcept", "l'outil produit un fichier de pré-saisie ; l'import est une action humaine.")
req(doc, "ENF-3", "Confidentialité / conformité", "données clients sensibles ; respect des politiques internes ; pas d'exfiltration vers des services externes non approuvés.")
req(doc, "ENF-4", "Robustesse", "erreurs renvoyées proprement (jamais de trace technique au client) ; un document en échec est signalé, pas masqué.")
req(doc, "ENF-5", "Traçabilité", "chaque étape est journalisée (fichier, client, statut, résultat) côté fonction.")
req(doc, "ENF-6", "Performance", "traitement unitaire par document, adapté au volume quotidien ; conversion d'image seulement si nécessaire.")
req(doc, "ENF-7", "Langue", "interface, colonnes et rapports en français.")
req(doc, "ENF-8", "Maintenabilité", "le comportement d'extraction/résolution est centralisé (prompts + schéma) ; le modèle IA est paramétrable ; le master data se met à jour par simple ré-export.")

# ============ 9. SCENARIO ============
doc.add_heading("9. Scénario nominal et cas limites", level=1)
doc.add_heading("9.1 Scénario nominal", level=2)
figure(doc, FLOW, "Figure 1 — Cycle de vie d'une commande. Bleu = automatique (l'outil) ; "
       "orange = action humaine (gestionnaire) ; vert = début/fin ; jaune = décision.")
para(doc, "Étapes :")
for i, s in enumerate([
    "Réception : dépôt automatique des pièces jointes de certains expéditeurs, puis tri humain.",
    "Le gestionnaire dépose chaque commande à traiter dans « Commandes-PDF ».",
    "Déclenchement du lot (quotidien à 15 h ou manuel).",
    "Lecture de chaque document par l'IA (client, référence, date, produits).",
    "Validation de format ; commande illisible/incomplète → « A-revoir ».",
    "Résolution master data : Clé 1 + confirmation/correction des SKU.",
    "Validation stricte ; client introuvable ou produit hors catalogue → « A-revoir ».",
    "Commande OK : PDF renommé et déplacé dans « Commandes-Done », ligne ajoutée au lot.",
    "Génération de l'Excel CCI consolidé dans « Uppload-CCI » (incertitudes surlignées jaune) + rapport « A-revoir ».",
    "Le gestionnaire vérifie TOUT l'Excel, reprend les « A-revoir », puis saisit dans ProConcept pour créer les CCI.",
], start=1):
    bullet(doc, [(f"{i}. ", True), (s, False)], style="List Number" if False else "List Bullet")

doc.add_heading("9.2 Cas limites", level=2)
table(doc,
      ["Situation", "Comportement de l'outil"],
      [
        ["Document illisible / corrompu", "Signalé et routé vers « A-revoir » ; jamais reconstitué depuis un document TRB."],
        ["Champ commande manquant (client / référence / date)", "Rejet en « A-revoir » avec le motif dans le rapport d'erreurs."],
        ["SKU client faux ou absent", "Corrigé via la désignation et le catalogue du client (n'est pas un motif de rejet)."],
        ["SKU ambigu (produits proches)", "SKU proposé mais surligné jaune dans l'Excel : à confirmer par l'humain."],
        ["Client introuvable dans le master data", "Rejet en « A-revoir » ; souvent un nouveau client à créer et un référentiel à compléter."],
        ["Produit hors catalogue du client", "Rejet en « A-revoir » ; produit jamais acheté par ce client jusqu'ici."],
        ["Clients homonymes (ex. TRB Thailand / Vietnam)", "L'adresse de livraison sert à départager le bon client."],
      ],
      widths=[2.5, 4.0])

# ============ 10. CRITERES ============
doc.add_heading("10. Critères d'acceptation", level=1)
table(doc,
      ["Critère", "Cible"],
      [
        ["CA-1 — Nom du document traité renommé « Client - JJ-MM-AAAA »", "100 %"],
        ["CA-2 — Écritures directes dans ProConcept", "0"],
        ["CA-3 — E-mails envoyés au client par l'outil", "0"],
        ["CA-4 — Excel CCI unique, horodaté, sans écrasement, par lot", "100 %"],
        ["CA-5 — Commande sans Clé 1 ou avec produit hors catalogue incluse dans l'Excel", "0 (routée en « A-revoir »)"],
        ["CA-6 — SKU ou Clé 1 inventés (absents du master data)", "0 (garde-fou)"],
        ["CA-7 — SKU ambigu / Clé 1 vide surlignés en jaune", "100 %"],
        ["CA-8 — Commandes rejetées listées dans le rapport « A-revoir »", "100 %"],
        ["CA-9 — Formats lus (PDF natif, scanné, PNG, JPEG, TIFF)", "Tous (jeu de tests)"],
        ["CA-10 — Un document en échec bloque le reste du lot", "0 (isolation des échecs)"],
      ],
      widths=[4.6, 1.9])

# ============ 11. HYPOTHESES ============
doc.add_heading("11. Hypothèses et dépendances", level=1)
bullet(doc, "Le master data (clients + catalogue), dérivé de l'historique des ventes, est disponible et rafraîchi périodiquement.")
bullet(doc, "Les dossiers SharePoint (Commandes-PDF, Uppload-CCI, Commandes-Done, A-revoir) existent et sont accessibles à l'outil.")
bullet(doc, "Une clé d'accès à l'IA (Claude) est configurée côté serveur.")
bullet(doc, "Les commandes arrivent par e-mail avec pièce jointe ; un gestionnaire assure le tri de réception.")
bullet(doc, "L'import dans ProConcept reste manuel (ou fera l'objet d'un lot ultérieur).")

# ============ 12. HORS PERIMETRE / ROADMAP ============
doc.add_heading("12. Hors périmètre et évolutions prévues", level=1)
para(doc, "Volontairement hors périmètre à ce jour, et pistes d'évolution :")
bullet(doc, [("Tri des e-mails plus intelligent", True), (" : détection sémantique des commandes pour "
             "réduire le tri humain — sachant qu'une part humaine restera nécessaire (l'automatisation ne captera pas tout).", False)])
bullet(doc, [("Brouillon de confirmation client", True), (" : préparation (sans envoi) d'une réponse au client.", False)])
bullet(doc, [("Fichiers de chargement ProConcept", True), (" : génération d'un onglet « Load » et/ou d'un XML pour accélérer l'import.", False)])
bullet(doc, [("Détection des anciens codes", True), (" : signalement des codes remplacés (ex. passage Tyvek) avec proposition du code actuel.", False)])
bullet(doc, [("Externalisation du master data", True), (" : lecture directe depuis SharePoint/ERP plutôt qu'embarqué avec l'outil.", False)])
bullet(doc, [("Tableau de bord de confiance", True), (" : suivi des taux de résolution et des motifs de rejet.", False)])

# ============ 13. GLOSSAIRE ============
doc.add_heading("13. Glossaire", level=1)
table(doc,
      ["Terme", "Définition"],
      [
        ["CCI", "Commande client interne : la commande saisie dans ProConcept à partir du bon de commande."],
        ["Clé 1", "Numéro client ProConcept (7 chiffres), clé de résolution du client. Retrouvé par l'IA à partir du nom."],
        ["SKU", "Code article TRB à 4 chiffres (« Référence principale » dans l'historique des ventes)."],
        ["Désignation", "Nom du produit tel qu'écrit sur la commande ; sert à retrouver le bon SKU."],
        ["Master data", "Référentiel de résolution : table clients (Clé 1) + catalogue par client (SKU + désignation)."],
        ["Catalogue (d'un client)", "Liste des produits (SKU + désignation) que ce client a réellement achetés."],
        ["Référence partenaire", "Numéro de la commande côté client (référence du bon de commande)."],
        ["A-revoir", "Statut de rejet : commande à traiter manuellement (illisible, client introuvable, produit hors catalogue)."],
        ["Statut SKU", "correct / corrigé / ambigu (à confirmer, surligné jaune) / inconnu (rejet)."],
        ["Pré-saisie", "Excel CCI intermédiaire à vérifier avant import dans ProConcept."],
        ["ProConcept", "ERP cible dans lequel les commandes (CCI) sont saisies."],
        ["Commandes-PDF / Uppload-CCI / Commandes-Done / A-revoir", "Dossiers SharePoint : entrée · sortie (Excel CCI) · archive des traités · rejetés."],
        ["Humain dans la boucle", "Validation humaine obligatoire avant toute action sortante (import ProConcept)."],
      ],
      widths=[2.2, 4.3])

doc.save(OUT)
print("OK:", OUT)
