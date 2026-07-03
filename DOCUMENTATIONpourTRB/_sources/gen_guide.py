# -*- coding: utf-8 -*-
"""Guide utilisateur — Outil CCI (TRB Chemedica)."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docgen_common import *
from docx.enum.text import WD_ALIGN_PARAGRAPH

FLOW = os.path.join(os.path.dirname(os.path.abspath(__file__)), "flowchart_cci.png")
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), r"Guide_utilisateur_Outil_CCI_v1.docx")

doc = new_doc()

# ============ PAGE DE TITRE ============
title_block(doc, [
    ("GUIDE UTILISATEUR", {"bold": True, "color": TRB_BLUE, "size": 24, "space_after": 2, "space_before": 40}),
    ("Outil CCI — Traitement automatique des commandes clients", {"bold": True, "color": DARK, "size": 15, "space_after": 2}),
    ("Mode d'emploi et procédure de vérification avant saisie dans ProConcept",
     {"italic": True, "color": GREY, "size": 12, "space_after": 16}),
])
info = [
    ("Pour", "Gestionnaires / spécialistes Supply Chain"),
    ("Organisation", "TRB Chemedica International SA"),
    ("Version", "1.0  —  3 juillet 2026"),
]
kv_block(doc, info)
para(doc, "")
callout(doc,
        "L'outil prépare le travail et signale ses doutes ; vous gardez la main. En particulier, "
        "c'est vous qui vérifiez l'Excel de pré-saisie et qui saisissez les commandes dans ProConcept. "
        "L'outil n'écrit jamais dans l'ERP et n'envoie aucun e-mail au client.",
        label="L'essentiel :", fill=INFO_FILL, edge="548235")

doc.add_page_break()

# ============ SOMMAIRE ============
doc.add_heading("Sommaire", level=1)
sommaire(doc, [
    ("Vue d'ensemble du processus", False),
    ("Partie A — Guide d'utilisation", False),
    ("A.1 En bref  ·  A.2 Avant de commencer  ·  A.3 Déposer les commandes", True),
    ("A.4 Ce qui se passe automatiquement  ·  A.5 Où retrouver les résultats  ·  A.6 Comprendre l'Excel CCI", True),
    ("Partie B — Procédure de vérification (avant import ProConcept)", False),
    ("B.0 Ouvrir  ·  B.1 Client & Clé 1  ·  B.2 SKU ★  ·  B.3 Quantités  ·  B.4 Référence & date", True),
    ("B.5 Commandes « A-revoir »  ·  B.6 Saisir dans ProConcept", True),
    ("Partie C — Aide-mémoire (à cocher pour chaque lot)", False),
    ("Partie D — Que faire si… (cas particuliers)", False),
    ("Rappels", False),
])

# ============ VUE D'ENSEMBLE ============
doc.add_heading("Vue d'ensemble du processus", level=1)
para(doc, "L'outil traite vos commandes par lot : chaque jour à 15 h (et à la demande), il lit tous "
     "les bons de commande déposés dans le dossier « Commandes-PDF », retrouve le numéro client "
     "(« Clé 1 ») et confirme/corrige les codes article (SKU), puis réunit toutes les commandes "
     "valides dans UN SEUL fichier Excel de pré-saisie déposé dans « Uppload-CCI ». Les commandes "
     "qu'il n'a pas pu traiter partent dans « A-revoir » avec un rapport expliquant pourquoi. "
     "Les documents traités sont archivés (renommés) dans « Commandes-Done ». "
     "Ensuite, vous vérifiez l'Excel et vous saisissez les commandes dans ProConcept. "
     "Le schéma ci-dessous situe chaque étape et, surtout, à quel moment votre intervention est requise.")
figure(doc, FLOW, "Cycle de vie d'une commande. Bleu = fait par l'outil (automatique) ; "
       "orange = fait par vous (le gestionnaire) ; vert = début/fin ; jaune = décision de l'outil.")

# ============ PARTIE A ============
doc.add_heading("Partie A — Guide d'utilisation", level=1)

doc.add_heading("A.1  En bref", level=2)
para(doc, "L'outil lit vos bons de commande, retrouve le client et les codes article à partir de "
     "l'historique des ventes, et regroupe toutes les commandes du lot dans un Excel CCI "
     "(1 ligne par commande). Il surligne en jaune ce dont il n'est pas sûr. Vous vérifiez cet "
     "Excel, vous reprenez les commandes du dossier « A-revoir », puis vous saisissez le tout "
     "dans ProConcept.")
callout(doc,
        "Ce que l'outil NE fait JAMAIS : il n'écrit jamais dans ProConcept (il produit un fichier à "
        "importer) et n'envoie aucun e-mail aux clients. La vérification finale et la saisie ERP "
        "restent votre responsabilité.",
        label="À retenir :")

doc.add_heading("A.2  Avant de commencer", level=2)
bullet(doc, "Le master data (l'historique des ventes qui contient les clients « Clé 1 » et les catalogues) doit être présent et à jour.")
bullet(doc, "Les quatre dossiers SharePoint existent : « Commandes-PDF » (entrée), « Uppload-CCI » (sortie), « Commandes-Done » (archive), « A-revoir » (rejetés).")
bullet(doc, "Les commandes arrivent par e-mail, en pièce jointe (PDF ou image, y compris scannée ou manuscrite).")

doc.add_heading("A.3  Déposer les commandes (réception & tri)", level=2)
para(doc, "Une règle de messagerie dépose automatiquement les pièces jointes de certains expéditeurs "
     "connus. Cette règle est un point de départ : ", space_after=3)
bullet(doc, [("Elle ne capte pas tout", True), (" : commande écrite dans le corps du mail, expéditeur "
             "non listé, pièce jointe qui n'est pas une commande, format inhabituel…", False)])
bullet(doc, [("Votre rôle : ", True), ("vérifier les documents reçus et déposer chaque vraie commande à "
             "traiter dans le dossier « Commandes-PDF ». Ce tri humain est indispensable — l'outil ne "
             "traite que ce qui se trouve dans ce dossier au moment du lot.", False)])

doc.add_heading("A.4  Ce qui se passe automatiquement", level=2)
para(doc, "Au déclenchement du lot (15 h ou lancement manuel), pour chaque document de « Commandes-PDF » :")
bullet(doc, "Lecture du document par l'IA : nom du client, référence partenaire (n° de commande), date de livraison, et par ligne désignation / SKU / quantité.")
bullet(doc, "Contrôle de lisibilité : un document illisible ou incomplet part directement en « A-revoir ».")
bullet(doc, "Croisement avec le master data : recherche du numéro client (Clé 1) et confirmation/correction de chaque SKU via le catalogue du client.")
bullet(doc, "Décision : client introuvable ou produit hors catalogue → « A-revoir » ; sinon la commande est retenue.")
bullet(doc, "Archivage : le PDF traité est renommé « Client - JJ-MM-AAAA » et déplacé dans « Commandes-Done ».")
bullet(doc, "Consolidation : toutes les commandes valides du lot sont réunies dans UN Excel CCI horodaté dans « Uppload-CCI », les points incertains surlignés en jaune.")
bullet(doc, "Reporting : un Excel liste les commandes parties en « A-revoir » avec le motif.")

doc.add_heading("A.5  Où retrouver les résultats", level=2)
table(doc,
      ["Élément", "Où le trouver"],
      [
        ["Excel CCI de pré-saisie", "SharePoint › « Uppload-CCI » — UN fichier « CCI-Lot-AAAAMMJJ-HHMMSS.xlsx » regroupant toutes les commandes du lot"],
        ["Documents traités (archive)", "SharePoint › « Commandes-Done » — PDF renommés « Client - JJ-MM-AAAA »"],
        ["Commandes à reprendre", "SharePoint › « A-revoir » — PDF renommés + un Excel « rapport d'erreurs »"],
        ["Motif d'un rejet", "Excel « rapport d'erreurs » : colonnes Nom du fichier · Date de l'erreur · Observation"],
      ],
      widths=[2.1, 4.4])

doc.add_heading("A.6  Comprendre l'Excel CCI (et le jaune)", level=2)
para(doc, "L'Excel CCI comporte une ligne par commande. Ses colonnes :")
table(doc,
      ["Colonne", "Contenu"],
      [
        ["Nom du client", "Le nom lu sur la commande"],
        ["Clé 1", "Le numéro client trouvé dans le master data"],
        ["Référence partenaire", "Le numéro de commande du client"],
        ["Date de livraison souhaitée", "La date lue sur la commande"],
        ["SKU 1 / Quantité 1, SKU 2 / Quantité 2, …", "Une paire de colonnes par produit de la commande (SKU confirmé/corrigé + quantité)"],
        ["Nom du fichier", "Le nom composé « Client - JJ-MM-AAAA » pour retrouver le document d'origine"],
      ],
      widths=[2.5, 4.0])
callout(doc,
        "Le surlignage jaune signale deux choses, et seulement celles-là : un SKU dont l'outil n'est "
        "pas sûr (il a hésité entre plusieurs produits proches), et une Clé 1 vide (anormale). "
        "Le jaune est une aide au repérage — l'absence de jaune ne garantit pas que tout est juste. "
        "Vous restez responsable de vérifier l'ensemble.",
        label="Le jaune = « à vérifier » :")

# ============ PARTIE B ============
doc.add_heading("Partie B — Procédure de vérification (avant import ProConcept)", level=1)
callout(doc,
        "Principe : vous êtes le contrôle qualité. Effectuez les vérifications ci-dessous sur l'Excel "
        "CCI AVANT toute saisie dans ProConcept. Vérifiez TOUTES les lignes, pas seulement les cellules "
        "jaunes. En cas de doute sur une ligne, ne la saisissez pas : vérifiez d'abord.")

doc.add_heading("B.0  Ouvrir les éléments", level=2)
bullet(doc, "Ouvrir l'Excel CCI du lot — SharePoint › « Uppload-CCI ».")
bullet(doc, "Garder sous la main les documents d'origine (dossier « Commandes-Done ») et le rapport « A-revoir ».")

def control(doc, code, title, verifier, comment, sinon):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(4)
    r = p.add_run(f"{code} — {title}"); r.bold = True; r.font.color.rgb = TRB_BLUE
    bullet(doc, [("À vérifier : ", True), (verifier, False)])
    bullet(doc, [("Comment : ", True), (comment, False)])
    bullet(doc, [("Si problème : ", True), (sinon, False)])

control(doc, "B.1", "Nom du client et Clé 1",
        "chaque commande a un numéro client (Clé 1) renseigné, cohérent avec le nom lu.",
        "lire la colonne « Clé 1 » et « Nom du client » ; une Clé 1 vide est surlignée jaune.",
        "si la Clé 1 est vide : le client n'a pas été retrouvé — identifiez-le dans ProConcept (ou créez-le s'il est nouveau) et reportez son numéro. Vérifiez qu'il ne s'agit pas d'un homonyme.")
control(doc, "B.2", "SKU (code article)  ★ point critique",
        "chaque SKU est cohérent avec la désignation du produit commandé.",
        "contrôler en priorité les SKU surlignés jaune (l'outil hésite) ; recouper avec le document d'origine et la fiche article. Attention aux variantes (marché, conditionnement).",
        "si un SKU semble faux : corrigez-le d'après la fiche article. Un produit vraiment hors catalogue aura fait partir la commande en « A-revoir ».")
control(doc, "B.3", "Quantités",
        "la quantité de chaque ligne correspond au bon de commande.",
        "comparer chaque « Quantité i » avec le document d'origine ; vérifier les multiples de carton plein et les éventuels échantillons gratuits.",
        "si une quantité ne correspond pas : corrigez-la et, si nécessaire, contactez le client.")
control(doc, "B.4", "Référence partenaire et date de livraison",
        "le numéro de commande client et la date de livraison sont corrects.",
        "comparer les colonnes « Référence partenaire » et « Date de livraison souhaitée » au document d'origine.",
        "compléter ou corriger si l'outil a mal lu (surtout sur les documents manuscrits ou scannés).")

doc.add_heading("B.5  Commandes du dossier « A-revoir »", level=2)
para(doc, "Ces commandes n'ont pas été traitées automatiquement. Ouvrez le rapport d'erreurs (colonne "
     "« Observation ») pour connaître le motif, puis traitez-les à la main :")
table(doc,
      ["Motif (Observation)", "Que faire"],
      [
        ["Document illisible / incomplet", "Rouvrir le document d'origine ; s'il est vraiment illisible, redemander un document lisible au client. Saisir la commande manuellement."],
        ["Client introuvable", "Souvent un nouveau client : le créer dans ProConcept, puis saisir la commande. Penser à mettre le master data à jour."],
        ["Produit hors catalogue", "Produit jamais acheté par ce client jusqu'ici : vérifier le SKU via la fiche article et saisir manuellement."],
      ],
      widths=[2.3, 4.2])

doc.add_heading("B.6  Saisir dans ProConcept", level=2)
bullet(doc, "Une fois l'Excel CCI vérifié (et corrigé si besoin), saisir / importer les commandes dans ProConcept pour créer les CCI.")
bullet(doc, "Des modifications restent possibles à ce stade, directement dans ProConcept.")
bullet(doc, "Conserver l'Excel CCI et les documents d'origine (« Commandes-Done ») comme trace.")

# ============ PARTIE C ============
doc.add_heading("Partie C — Aide-mémoire (à cocher pour chaque lot)", level=1)
table(doc,
      ["#", "Contrôle", "Fait"],
      [
        ["1", "Toutes les commandes du lot ont une Clé 1 renseignée et correcte", "☐"],
        ["2", "SKU cohérents avec la désignation — surlignés jaune vérifiés en priorité  ★", "☐"],
        ["3", "Quantités correctes (carton plein, échantillons)", "☐"],
        ["4", "Référence partenaire et date de livraison correctes", "☐"],
        ["5", "Commandes du dossier « A-revoir » reprises et traitées", "☐"],
        ["6", "Excel CCI vérifié en entier (pas seulement le jaune)", "☐"],
        ["7", "Commandes saisies dans ProConcept (CCI créées)", "☐"],
        ["8", "Excel CCI et documents d'origine conservés", "☐"],
      ],
      widths=[0.5, 5.2, 0.7])

# ============ PARTIE D ============
doc.add_heading("Partie D — Que faire si… (cas particuliers)", level=1)
table(doc,
      ["Situation", "Que faire"],
      [
        ["Une commande n'apparaît ni dans l'Excel CCI ni dans « A-revoir »", "Vérifier qu'elle a bien été déposée dans « Commandes-PDF » avant le lot ; sinon la déposer et relancer / attendre le prochain lot."],
        ["Un SKU est surligné jaune", "L'outil n'est pas sûr : comparer avec le document d'origine et la fiche article, puis confirmer ou corriger."],
        ["La Clé 1 est vide (jaune)", "Client non retrouvé : l'identifier ou le créer dans ProConcept ; mettre le master data à jour pour la prochaine fois."],
        ["Commande dans « A-revoir »", "Lire l'« Observation » du rapport d'erreurs et traiter manuellement selon le motif (§ B.5)."],
        ["Bon de commande illisible", "Vérifier la pièce jointe ; si vraiment illisible, redemander un document lisible au client."],
        ["Beaucoup de rejets pour « client introuvable » / « hors catalogue »", "Signe d'un master data à rafraîchir : demander un nouvel export de l'historique des ventes."],
      ],
      widths=[2.5, 4.0])

# ============ RAPPELS ============
doc.add_heading("Rappels", level=1)
bullet(doc, "L'outil n'écrit jamais dans ProConcept : la saisie/import est votre action, après vérification.")
bullet(doc, "L'outil n'envoie aucun e-mail au client.")
bullet(doc, "Le surlignage jaune aide, mais ne remplace pas votre contrôle : vérifiez toutes les lignes.")
bullet(doc, "Un master data à jour = moins de « A-revoir » et des SKU mieux résolus. Signalez les nouveaux clients / produits.")

doc.save(OUT)
print("OK:", OUT)
