#!/usr/bin/env python3
"""Construit la master data de SUBSTITUTION depuis demo/faked_orders.json.

Écrit cci-function/app/master_data.xlsx (2 feuilles, même format que la vraie) :
  - clients   : customer_code | nom_client       — tous les clients SAUF les fail_client
  - catalogue : customer_code | sku | designation — les produits de chaque commande,
                SAUF la ligne « intrus » des fail_product (→ produit hors catalogue)

La vraie master data doit avoir été sauvegardée au préalable :
    cci-function/app/master_data.REEL.backup.xlsx

Avant d'écrire, le script VÉRIFIE les invariants qui rendent la démo déterministe
(voir docs/superpowers/specs/2026-07-03-demo-fake-orders-design.md). Il échoue
bruyamment si une commande `pass` n'est pas entièrement couverte, ou si une
omission de `fail` n'est pas en place.

Usage :
    python3 demo/build_demo_master_data.py          # vérifie + écrit
    python3 demo/build_demo_master_data.py --check   # vérifie seulement (n'écrit pas)
"""
import json
import os
import sys

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
JSON_PATH = os.path.join(HERE, "faked_orders.json")
MD_PATH = os.path.join(REPO, "cci-function", "app", "master_data.xlsx")
BACKUP_PATH = os.path.join(REPO, "cci-function", "app", "master_data.REEL.backup.xlsx")


def load_orders():
    with open(JSON_PATH, encoding="utf-8") as f:
        return json.load(f)


def derive(orders):
    """→ (clients:[(code,nom)], catalogue:[(code,sku,desig)]) selon les outcomes."""
    clients = []
    catalogue = []
    seen_cat = set()  # (code, sku) dédup
    for o in orders:
        if o["outcome"] == "fail_client":
            continue  # client absent de la master data → 422 client introuvable
        clients.append((o["customer_code"], o["client_name"]))
        intruder = o.get("intruder_index")
        for i, ln in enumerate(o["lines"]):
            if o["outcome"] == "fail_product" and i == intruder:
                continue  # ligne intrus omise → 422 produit hors catalogue
            key = (o["customer_code"], ln["sku"])
            if key in seen_cat:
                continue
            seen_cat.add(key)
            catalogue.append((o["customer_code"], ln["sku"], ln["designation"]))
    return clients, catalogue


def verify(orders, clients, catalogue):
    """Asserte les invariants qui garantissent le pass/fail de la démo."""
    names = {c[1] for c in clients}
    names_lower = {c[1].lower() for c in clients}
    codes = {c[0] for c in clients}
    cat_pairs = {}  # code -> set((sku, designation))
    for code, sku, des in catalogue:
        cat_pairs.setdefault(code, set()).add((sku, des))

    errors = []
    for o in orders:
        tag = f'{o["out"]} ({o["outcome"]})'
        if o["outcome"] == "pass":
            if o["client_name"] not in names:
                errors.append(f"[{tag}] client absent des clients : {o['client_name']}")
            pairs = cat_pairs.get(o["customer_code"], set())
            for ln in o["lines"]:
                if (ln["sku"], ln["designation"]) not in pairs:
                    errors.append(f"[{tag}] ligne hors catalogue : {ln['sku']} | {ln['designation']}")
        elif o["outcome"] == "fail_client":
            if o["client_name"] in names:
                errors.append(f"[{tag}] client PRÉSENT alors qu'il doit être absent")
            # robustesse anti-match flou : nom clairement distinct des clients présents
            if o["client_name"].lower() in names_lower:
                errors.append(f"[{tag}] nom identique (casse) à un client présent")
        elif o["outcome"] == "fail_product":
            if o["client_name"] not in names:
                errors.append(f"[{tag}] client absent alors qu'il doit être présent")
            pairs = cat_pairs.get(o["customer_code"], set())
            intr = o["lines"][o["intruder_index"]]
            if (intr["sku"], intr["designation"]) in pairs:
                errors.append(f"[{tag}] ligne intrus PRÉSENTE au catalogue : {intr['sku']} | {intr['designation']}")
            for i, ln in enumerate(o["lines"]):
                if i == o["intruder_index"]:
                    continue
                if (ln["sku"], ln["designation"]) not in pairs:
                    errors.append(f"[{tag}] ligne (non-intrus) hors catalogue : {ln['sku']} | {ln['designation']}")

    # format
    for code, nom in clients:
        if not (len(code) == 7 and code.isdigit()):
            errors.append(f"code client mal formé : {code}")
    for code, sku, des in catalogue:
        if not (len(sku) == 4 and sku.isdigit()):
            errors.append(f"sku mal formé : {sku} ({code})")
    if len(codes) != len(clients):
        errors.append("codes client non uniques")

    if errors:
        print("✗ INVARIANTS VIOLÉS :")
        for er in errors:
            print("   -", er)
        raise SystemExit(1)
    print(f"✓ invariants OK — {len(clients)} clients, {len(catalogue)} lignes catalogue")


def write_xlsx(clients, catalogue):
    if not os.path.exists(BACKUP_PATH):
        raise SystemExit(
            f"✗ backup de la vraie master data introuvable : {BACKUP_PATH}\n"
            "  Crée-le d'abord (cp master_data.xlsx master_data.REEL.backup.xlsx)."
        )
    wb = Workbook()
    ws_c = wb.active
    ws_c.title = "clients"
    ws_c.append(["customer_code", "nom_client"])
    for code, nom in clients:
        ws_c.append([code, nom])
    ws_cat = wb.create_sheet("catalogue")
    ws_cat.append(["customer_code", "sku", "designation"])
    for code, sku, des in catalogue:
        ws_cat.append([code, sku, des])
    # SKU / code en TEXTE (préserve les zéros de tête, ex. 0587)
    for ws, cols in ((ws_c, ["A"]), (ws_cat, ["A", "B"])):
        for col in cols:
            for cell in ws[col][1:]:
                cell.number_format = "@"
    wb.save(MD_PATH)
    print(f"✓ écrit {MD_PATH}")


def main():
    orders = load_orders()
    clients, catalogue = derive(orders)
    verify(orders, clients, catalogue)
    if "--check" in sys.argv:
        print("(--check : master_data.xlsx non modifié)")
        return
    write_xlsx(clients, catalogue)
    n_fc = sum(o["outcome"] == "fail_client" for o in orders)
    n_fp = sum(o["outcome"] == "fail_product" for o in orders)
    print(f"  → {n_fc} clients volontairement absents (fail_client)")
    print(f"  → {n_fp} commandes avec 1 produit hors catalogue (fail_product)")


if __name__ == "__main__":
    main()
